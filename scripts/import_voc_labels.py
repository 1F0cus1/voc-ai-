"""
Import VOC label taxonomy from Excel into MySQL/MariaDB.

Default source sheet: VOC标签确认于2026年5月20
Default target table: voc_label_taxonomy

Install dependency if needed:
    pip install openpyxl pymysql

Example:
    set DB_HOST=127.0.0.1
    set DB_PORT=3306
    set DB_USER=root
    set DB_PASSWORD=your_password
    set DB_NAME=your_database
    python scripts/import_voc_labels.py ^
      --excel "C:\\Users\\Admin\\Desktop\\客诉VOC项目\\客诉评价标签分类.xlsx" ^
      --label-version 2026-05-20
"""

from __future__ import annotations

import argparse
import os
import sys
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook


DEFAULT_EXCEL_PATH = r"C:\Users\Admin\Desktop\客诉VOC项目\客诉评价标签分类.xlsx"
DEFAULT_SHEET_NAME = "VOC标签确认于2026年5月20"
DEFAULT_TABLE_NAME = "voc_label_taxonomy"


@dataclass
class DbConfig:
    host: str
    port: int
    user: str
    password: str
    database: str
    charset: str = "utf8mb4"


def text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def get_db_config() -> DbConfig:
    missing = [
        name
        for name in ("DB_HOST", "DB_USER", "DB_PASSWORD", "DB_NAME")
        if not os.getenv(name)
    ]
    if missing:
        raise SystemExit(
            "Missing database environment variables: "
            + ", ".join(missing)
            + "\nRequired: DB_HOST, DB_PORT(optional), DB_USER, DB_PASSWORD, DB_NAME"
        )

    return DbConfig(
        host=os.environ["DB_HOST"],
        port=int(os.getenv("DB_PORT", "3306")),
        user=os.environ["DB_USER"],
        password=os.environ["DB_PASSWORD"],
        database=os.environ["DB_NAME"],
    )


def build_merged_lookup(ws) -> dict[tuple[int, int], Any]:
    lookup: dict[tuple[int, int], Any] = {}
    for merged_range in ws.merged_cells.ranges:
        top_left_value = ws.cell(merged_range.min_row, merged_range.min_col).value
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for col in range(merged_range.min_col, merged_range.max_col + 1):
                lookup[(row, col)] = top_left_value
    return lookup


def cell_value_with_merged(ws, row: int, col: int, merged_lookup: dict[tuple[int, int], Any]) -> Any:
    cell = ws.cell(row=row, column=col)
    if cell.value is not None:
        return cell.value
    return merged_lookup.get((row, col))


def choose_sheet(wb, requested_name: str):
    if requested_name in wb.sheetnames:
        return wb[requested_name]

    for ws in wb.worksheets:
        if "VOC" in ws.title:
            return ws

    raise SystemExit(
        f"Sheet not found: {requested_name}. Available sheets: {', '.join(wb.sheetnames)}"
    )


def load_labels(excel_path: str, sheet_name: str) -> list[dict[str, str]]:
    wb = load_workbook(excel_path, read_only=False, data_only=True)
    ws = choose_sheet(wb, sheet_name)
    merged_lookup = build_merged_lookup(ws)

    labels: list[dict[str, str]] = []
    seen: set[tuple[str, str, str, str]] = set()

    for row in range(2, ws.max_row + 1):
        item = {
            "level1_name": text(cell_value_with_merged(ws, row, 1, merged_lookup)),
            "level2_name": text(cell_value_with_merged(ws, row, 2, merged_lookup)),
            "level3_name": text(cell_value_with_merged(ws, row, 3, merged_lookup)),
            "level4_name": text(cell_value_with_merged(ws, row, 4, merged_lookup)),
            "definition": text(cell_value_with_merged(ws, row, 5, merged_lookup)),
            "keywords": text(cell_value_with_merged(ws, row, 6, merged_lookup)),
            "examples": text(cell_value_with_merged(ws, row, 7, merged_lookup)),
        }

        if not item["level4_name"]:
            continue

        key = (
            item["level1_name"],
            item["level2_name"],
            item["level3_name"],
            item["level4_name"],
        )
        if key in seen:
            continue

        missing_levels = [name for name in key if not name]
        if missing_levels:
            raise SystemExit(f"Row {row} has empty category level: {item}")

        seen.add(key)
        labels.append(item)

    if not labels:
        raise SystemExit("No labels found in Excel sheet.")

    return labels


def import_labels(
    db: DbConfig,
    table_name: str,
    label_version: str,
    labels: list[dict[str, str]],
    replace: bool,
) -> None:
    try:
        import pymysql
    except ImportError as exc:
        raise SystemExit("Missing dependency: pymysql. Install it with: pip install pymysql") from exc

    conn = pymysql.connect(
        host=db.host,
        port=db.port,
        user=db.user,
        password=db.password,
        database=db.database,
        charset=db.charset,
        autocommit=False,
    )

    insert_sql = f"""
        INSERT INTO `{table_name}` (
            label_version,
            level1_name,
            level2_name,
            level3_name,
            level4_name,
            definition,
            keywords,
            examples,
            enabled
        ) VALUES (
            %s, %s, %s, %s, %s, %s, %s, %s, 1
        )
    """

    rows = [
        (
            label_version,
            item["level1_name"],
            item["level2_name"],
            item["level3_name"],
            item["level4_name"],
            item["definition"] or None,
            item["keywords"] or None,
            item["examples"] or None,
        )
        for item in labels
    ]

    try:
        with conn.cursor() as cursor:
            if replace:
                cursor.execute(
                    f"DELETE FROM `{table_name}` WHERE label_version = %s",
                    (label_version,),
                )
            cursor.executemany(insert_sql, rows)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Import VOC label taxonomy into MySQL.")
    parser.add_argument("--excel", default=DEFAULT_EXCEL_PATH, help="Excel file path.")
    parser.add_argument("--sheet", default=DEFAULT_SHEET_NAME, help="Excel sheet name.")
    parser.add_argument(
        "--label-version",
        default="2026-05-20",
        help="Version value written to voc_label_taxonomy.label_version.",
    )
    parser.add_argument("--table", default=DEFAULT_TABLE_NAME, help="Target table name.")
    parser.add_argument(
        "--append",
        action="store_true",
        help="Append instead of deleting existing rows with the same label version first.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Only read and print summary; do not connect to database.",
    )
    return parser.parse_args()


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    args = parse_args()

    labels = load_labels(args.excel, args.sheet)
    print(f"Loaded {len(labels)} labels from Excel.")
    print("First 3 labels:")
    for item in labels[:3]:
        print(
            " - "
            + " / ".join(
                [
                    item["level1_name"],
                    item["level2_name"],
                    item["level3_name"],
                    item["level4_name"],
                ]
            )
        )

    if args.dry_run:
        print("Dry run only. Database was not changed.")
        return 0

    db = get_db_config()
    import_labels(
        db=db,
        table_name=args.table,
        label_version=args.label_version,
        labels=labels,
        replace=not args.append,
    )
    print(
        f"Imported {len(labels)} labels into {args.table}, "
        f"label_version={args.label_version}."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
